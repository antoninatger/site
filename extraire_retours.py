# -*- coding: utf-8 -*-
"""
extraire_retours.py — Choisit dans Retours.xlsx les retours à afficher sur le site.

Retours.xlsx (dossier « Web apps ») est rempli par le Hub, bouton « 💬 Retours → Excel ».
Ce script en tire `contenu/retours.yaml`, que construire_site.py lit (variable `retours`
dans les gabarits). Il est lancé automatiquement à chaque construction du site, quand
l'Excel est accessible (donc sur ton PC, pas sur GitHub : là-bas, c'est le retours.yaml
déjà publié qui sert).

Ce qui est affiché pour chaque retour :
  · le commentaire ;
  · prof : son nom et son établissement ;
  · élève : « Élève de 4e » (jamais de nom) et l'établissement ;
  · jeu : le nom du jeu (et le niveau joué), sans le prénom du joueur ;
  · présentiel ou distanciel (colonne « Présentiel / distanciel ») ;
  · l'intervention et la note sur 20 quand il y en a une.

Qui est retenu, dans l'Excel :
  · colonne « Sur le site » = oui → toujours ; = non → jamais ;
  · case vide → retenu si la note est d'au moins 16/20 (retours d'intervention) ou si le
    message est clairement positif (jeux), et si le commentaire fait au moins 30 caractères ;
  · colonne « Texte pour le site » remplie → c'est ce texte qui est publié (pour corriger
    une faute ou raccourcir), sinon l'« Opinion générale » (ou le message du jeu).

    python extraire_retours.py            # met à jour contenu/retours.yaml et affiche la liste
"""
import re, sys, unicodedata
from datetime import datetime
from pathlib import Path

ICI = Path(__file__).resolve().parent
EXCEL = ICI.parent.parent / "Retours.xlsx"          # Web apps/Retours.xlsx
SORTIE = ICI / "contenu" / "retours.yaml"
NOTE_MINI = 16
LONGUEUR_MINI = 30
LONGUEUR_MAXI = 260
MAX_RETOURS = 12

POSITIF = re.compile(r"\b(super|genial|interessant|merci|bravo|top|excellent|adore|aime|utile|"
                     r"passionnant|trop bien|tres bien|cool|parfait|clair|instructif|amusant|drole)", re.I)
NEGATIF = re.compile(r"\b(bug|erreur|confusion|probleme|faux|nul|ennuy|pas (clair|bien|interessant))", re.I)


def _sans_accents(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


def _propre(t):
    t = re.sub(r"\s+", " ", str(t or "")).strip().strip("—-– ").strip()
    if len(t) > LONGUEUR_MAXI:
        coupe = t[:LONGUEUR_MAXI]
        fin = max(coupe.rfind(". "), coupe.rfind("! "), coupe.rfind("? "))
        t = coupe[:fin + 1] if fin > 80 else coupe.rsplit(" ", 1)[0] + "…"
    return t


def _classe(c):
    c = str(c or "").strip()
    m = re.match(r"^(\d)\s*(?:e|eme|ème|°)", c, re.I)
    if m:
        return f"{m.group(1)}e"
    return c


def _note(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"(\d+(?:[.,]\d+)?)", str(v or ""))
    return float(m.group(1).replace(",", ".")) if m else None


def _date(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else ""


def _lignes(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    h = [str(x or "") for x in rows[0]]
    return [dict(zip(h, r)) for r in rows[1:] if any(r)]


def _choix(ligne):
    """'oui', 'non' ou '' d'après la colonne « Sur le site »."""
    v = _sans_accents(ligne.get("Sur le site")).strip()
    if v in ("oui", "o", "x", "1", "yes", "vrai"):
        return "oui"
    if v in ("non", "n", "0", "no", "faux"):
        return "non"
    return ""


def extraire():
    import openpyxl
    wb = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)
    retours, vus = [], set()

    def ajouter(r):
        cle = _sans_accents(r["texte"])[:80]
        if cle in vus:          # le même retour transféré deux fois
            return
        vus.add(cle)
        retours.append(r)

    for feuille, sorte in (("Retours profs", "prof"), ("Retours élèves", "eleve")):
        if feuille not in wb.sheetnames:
            continue
        for l in _lignes(wb[feuille]):
            choix = _choix(l)
            if choix == "non":
                continue
            texte = _propre(l.get("Texte pour le site") or l.get("Opinion générale"))
            if len(texte) < LONGUEUR_MINI:
                texte = _propre(l.get("Texte pour le site") or l.get("Points pertinents et clairs") or texte)
            note = _note(l.get("Note (/20)"))
            if choix != "oui" and (len(texte) < LONGUEUR_MINI or note is None or note < NOTE_MINI):
                continue
            if sorte == "prof":
                qui = " ".join(str(x).strip() for x in (l.get("Prénom"), l.get("Nom")) if x).strip() or "Enseignant·e"
            else:
                cl = _classe(l.get("Classe"))
                qui = f"Élève de {cl}" if cl else "Élève"
            ajouter({
                "type": sorte,
                "texte": texte,
                "qui": qui,
                "etablissement": str(l.get("Établissement") or "").strip(),
                "mode": str(l.get("Présentiel / distanciel") or "").strip(),
                "sujet": str(l.get("Intervention") or "").strip(),
                "note": (int(note) if note is not None and float(note).is_integer() else note),
                "date": _date(l.get("Date de l'intervention")) or _date(l.get("Reçu le")),
            })

    if "FakeMètre" in wb.sheetnames:
        for l in _lignes(wb["FakeMètre"]):
            choix = _choix(l)
            if choix == "non":
                continue
            objet = _sans_accents(l.get("Objet"))
            # Seuls les messages envoyés depuis le jeu (pas les échanges de mails autour du jeu)
            # (ou une partie arrivée sous l'objet générique du formulaire : elle a
            # son niveau et son score, 7 octobre 2026)
            fin_de_partie = re.search(r"fakemetre\s*[—–-]\s*(message de fin de partie|retour utilisateur)", objet) \
                or (str(l.get("Score") or "").strip() and str(l.get("Niveau") or "").strip())
            if choix != "oui" and not fin_de_partie:
                continue
            texte = _propre(l.get("Texte pour le site") or l.get("Message au créateur"))
            sa = _sans_accents(texte)
            if choix != "oui" and (len(texte) < LONGUEUR_MINI or not POSITIF.search(sa) or NEGATIF.search(sa)):
                continue
            niveau = str(l.get("Niveau") or "").strip()
            ajouter({
                "type": "jeu",
                "texte": texte,
                "qui": f"Partie de Fakemètre" + (f", niveau {niveau}" if niveau else ""),
                "etablissement": "",
                "mode": "",
                "sujet": "Fakemètre",
                "note": None,
                "date": _date(l.get("Reçu le")),
            })

    # Profs d'abord à date égale, puis les plus récents en tête
    rang = {"prof": 0, "eleve": 1, "jeu": 2}
    retours.sort(key=lambda r: (r["date"], -rang[r["type"]]), reverse=True)
    return retours[:MAX_RETOURS]


def ecrire(retours):
    import yaml
    entete = ("# Généré par extraire_retours.py à partir de Retours.xlsx — ne pas modifier ici :\n"
              "# pour retirer ou forcer un retour, colonne « Sur le site » (oui / non) dans l'Excel ;\n"
              "# pour corriger un texte, colonne « Texte pour le site ».\n")
    SORTIE.write_text(entete + yaml.safe_dump({"retours": retours}, allow_unicode=True, sort_keys=False, width=1000),
                      encoding="utf-8")


def mettre_a_jour(bavard=True):
    """Appelé par construire_site.py. Ne casse jamais la construction."""
    if not EXCEL.exists():
        return None
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        if bavard:
            print("Retours : openpyxl absent (pip install openpyxl) — retours.yaml non mis à jour.")
        return None
    try:
        r = extraire()
    except PermissionError:
        if bavard:
            print("Retours : Retours.xlsx est ouvert ailleurs — retours.yaml garde la version précédente.")
        return None
    except Exception as e:
        if bavard:
            print(f"Retours : lecture impossible ({e}) — retours.yaml garde la version précédente.")
        return None
    ecrire(r)
    if bavard:
        manque = sum(1 for x in r if x["type"] != "jeu" and not x["mode"])
        print(f"Retours : {len(r)} retour(s) retenu(s) pour le site"
              + (f", dont {manque} sans « Présentiel / distanciel » (à remplir dans Retours.xlsx)" if manque else "") + ".")
    return r


if __name__ == "__main__":
    r = mettre_a_jour()
    if r is None:
        print(f"Rien fait : {EXCEL} introuvable ou illisible.")
        sys.exit(1)
    for x in r:
        print(f"- [{x['type']}] {x['qui']} · {x['etablissement'] or x['sujet']} · {x['mode'] or '?'} · "
              f"{x['note'] if x['note'] is not None else ''}\n    {x['texte']}")
